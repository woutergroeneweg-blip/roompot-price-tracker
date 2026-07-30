#!/usr/bin/env python3
"""Append a price check to the Roompot tracker workbook.

This is the small, scriptable core of the tracker: given the total price
observed for the fixed booking, it appends a new row to the ``Price History``
sheet of ``prices.xlsx``, deriving the ``change`` and ``lowest_total_eur_so_far``
columns automatically, and keeps the trend chart pointed at the full data range.

It intentionally does NOT scrape the Roompot page itself: the live total is
rendered client-side, so the Cursor Automation reads it and passes it in via
``--total``. Everything spreadsheet-related lives here so it is testable.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl

SHEET = "Price History"
TZ = ZoneInfo("Europe/Amsterdam")
HEADERS = [
    "checked_at",
    "total_eur",
    "stay_eur",
    "tourist_tax_eur",
    "status",
    "change",
    "lowest_total_eur_so_far",
    "notes",
]


def _data_rows(ws):
    """Return existing data rows (list of tuples) below the header."""
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            continue
        rows.append(row)
    return rows


def _compute_change(prev_total, total, status):
    if status != "available" or total is None:
        return "unavailable"
    if prev_total is None:
        return "first"
    delta = int(total) - int(prev_total)
    return f"+{delta}" if delta > 0 else str(delta)


def _compute_lowest(prev_lowest, total, status):
    if status != "available" or total is None:
        return prev_lowest
    if prev_lowest is None:
        return int(total)
    return min(int(prev_lowest), int(total))


def _extend_chart(ws, last_row):
    for chart in ws._charts:
        for s in chart.series:
            for ref in (getattr(s, "val", None), getattr(s, "cat", None)):
                if ref is None:
                    continue
                num = getattr(ref, "numRef", None) or getattr(ref, "strRef", None)
                if num is None or not num.f:
                    continue
                col = num.f.split("$")[1] if "$" in num.f else None
                if col:
                    num.f = f"'{SHEET}'!${col}$2:${col}${last_row}"


def append_check(path, total, stay=None, tax=None, status="available",
                 notes="", timestamp=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]

    existing = _data_rows(ws)
    prev_available_total = None
    for r in reversed(existing):
        if r[4] == "available" and r[1] is not None:
            prev_available_total = r[1]
            break
    prev_lowest = existing[-1][6] if existing else None

    change = _compute_change(prev_available_total, total, status)
    lowest = _compute_lowest(prev_lowest, total, status)
    ts = timestamp or datetime.now(TZ).isoformat(timespec="seconds")

    new_row = [ts, total, stay, tax, status, change, lowest, notes]
    ws.append(new_row)
    _extend_chart(ws, ws.max_row)

    wb.save(path)
    return dict(zip(HEADERS, new_row))


def main():
    p = argparse.ArgumentParser(description="Append a price check to the tracker workbook.")
    p.add_argument("--file", default="prices.xlsx", help="Path to the workbook (default: prices.xlsx)")
    p.add_argument("--total", type=int, required=True, help="Observed total price in EUR")
    p.add_argument("--stay", type=int, default=None, help="Stay portion in EUR")
    p.add_argument("--tax", type=int, default=None, help="Tourist tax in EUR")
    p.add_argument("--status", default="available", choices=["available", "unavailable"])
    p.add_argument("--notes", default="", help="Free-text notes for the row")
    p.add_argument("--timestamp", default=None, help="Override ISO timestamp (default: now, Europe/Amsterdam)")
    args = p.parse_args()

    row = append_check(
        args.file, args.total, args.stay, args.tax, args.status, args.notes, args.timestamp
    )
    print("Appended row:")
    for k, v in row.items():
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
